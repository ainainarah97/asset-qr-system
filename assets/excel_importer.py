from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.db import transaction

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .models import (
    Premise,
    Block,
    Component,
    Level,
    Room,
    RoomFunctionCode,
    DynamicQRCode,
)

import uuid

# =========================================================
# DA6 HEADERS
# =========================================================

ROOM_HEADERS = [
    "BIL",
    "NAMA ARAS",
    "KOD ARAS",
    "LUAS ARAS (m2)",
    "KOD BLOK",
    "NAMA RUANG",
    "KOD RUANG",
    "FUNGSI RUANG",
    "KOD FUNGSI",
    "LUAS RUANG (m2)",
    "TINGGI (m)",
    "TEG RUANG",
    "CATATAN",
]


# =========================================================
# DAKKOMPONEN REQUIRED HEADERS
# =========================================================
#
# IMPORTANT:
#
# The real DA6 workbook uses:
#
#   LOKASI
#   NAMA KOMPONEN
#   KOD BIDANG KEJURUTERAAN
#   SISTEM
#   NO. PEROLEHAN ...
#
# The header helper below automatically recognises
# variations such as:
#
#   NO. PEROLEHAN
#   NO PEROLEHAN
#   NO. PEROLEHAN (1GFMA...)
#   NO ID KOMPONEN
#
# as:
#
#   NO. PEROLEHAN
#
# =========================================================

COMPONENT_REQUIRED_HEADERS = {
    "LOKASI",
    "NAMA KOMPONEN",
    "KOD BIDANG KEJURUTERAAN",
    "SISTEM",
    "NO. PEROLEHAN",
}


# =========================================================
# PREMISE WORKBOOK HEADERS
# =========================================================

PREMISE_REQUIRED_HEADERS = {
    "DPA NUMBER",
    "PREMISE NAME",
    "ADDRESS",
    "STATE",
}


# =========================================================
# GENERAL HELPERS
# =========================================================

def open_excel_workbook(file_path):
    """
    Open an XLSX or XLSM workbook safely.

    openpyxl does not execute VBA macros.
    """

    file_name = str(
        file_path
    ).lower()

    keep_vba = (
        file_name.endswith(
            ".xlsm"
        )
    )

    return load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
        keep_vba=keep_vba,
    )


def clean_text(value):
    if value is None:
        return ""

    if (
        isinstance(value, float)
        and value.is_integer()
    ):
        return str(
            int(value)
        )

    return str(value).strip()


def meaningful(value):
    text = clean_text(
        value
    )

    return text not in {
        "",
        "-",
    }


def normalize_header(value):
    """
    Normalise Excel headers so variations
    in spaces, line breaks and underscores
    can be compared reliably.
    """

    text = clean_text(
        value
    ).upper()

    # Non-breaking space
    text = text.replace(
        "\xa0",
        " ",
    )

    text = text.replace(
        "_",
        " ",
    )

    text = text.replace(
        "\n",
        " ",
    )

    text = text.replace(
        "\r",
        " ",
    )

    text = text.replace(
        "\t",
        " ",
    )

    while "  " in text:
        text = text.replace(
            "  ",
            " ",
        )

    return text.strip()


def clean_headers(values):
    headers = [
        clean_text(value)
        for value in values
    ]

    while (
        headers
        and headers[-1] == ""
    ):
        headers.pop()

    return headers


def find_header_row(
    worksheet,
    expected_headers,
):
    """
    Find an exact header row.

    Used for DAKRuang because its
    official input columns are stable.
    """

    maximum_row = min(
        worksheet.max_row or 1,
        30,
    )

    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=1,
            max_row=maximum_row,
            values_only=True,
        ),
        start=1,
    ):

        headers = clean_headers(
            row
        )

        if (
            headers
            == expected_headers
        ):
            return (
                row_number,
                headers,
            )

    return (
        None,
        [],
    )


def build_first_header_map(row):
    """
    Build a header -> column index map.

    Only the FIRST occurrence of a header
    is stored because DAKKomponen contains
    duplicate headers later in the sheet.

    This function also creates aliases for
    different DA6 component header formats.
    """

    header_map = {}

    for index, value in enumerate(
        row
    ):

        header = normalize_header(
            value
        )

        if not header:
            continue

        # -------------------------------------------------
        # Store original normalised header
        # -------------------------------------------------

        if header not in header_map:
            header_map[
                header
            ] = index

        # -------------------------------------------------
        # DAKKOMPONEN: NO. PEROLEHAN aliases
        # -------------------------------------------------
        #
        # Examples accepted:
        #
        # NO. PEROLEHAN
        # NO PEROLEHAN
        # NO. PEROLEHAN (1GFMA...)
        # NO PEROLEHAN (1GFMA...)
        # NO ID KOMPONEN
        #
        # -------------------------------------------------

        if (
            header.startswith(
                "NO. PEROLEHAN"
            )
            or header.startswith(
                "NO PEROLEHAN"
            )
            or header.startswith(
                "NO.  PEROLEHAN"
            )
            or header == (
                "NO ID KOMPONEN"
            )
            or header == (
                "NO. ID KOMPONEN"
            )
        ):

            if (
                "NO. PEROLEHAN"
                not in header_map
            ):
                header_map[
                    "NO. PEROLEHAN"
                ] = index

        # -------------------------------------------------
        # LOCATION alias
        # -------------------------------------------------
        #
        # New DA6 = LOKASI
        # Older version may use KOD LOKASI
        #
        # -------------------------------------------------

        if header == "KOD LOKASI":

            if (
                "LOKASI"
                not in header_map
            ):
                header_map[
                    "LOKASI"
                ] = index

        # -------------------------------------------------
        # SUBSYSTEM aliases
        # -------------------------------------------------

        if header == "SUBSISTEM":

            if (
                "SUB SISTEM"
                not in header_map
            ):
                header_map[
                    "SUB SISTEM"
                ] = index

        if header == "KOD SUB SISTEM":

            if (
                "KOD SUBSISTEM"
                not in header_map
            ):
                header_map[
                    "KOD SUBSISTEM"
                ] = index

    return header_map


def find_required_header_row(
    worksheet,
    required_headers,
    max_search_rows=50,
):
    """
    Find a row containing all required headers.

    Header values are normalised and aliases
    from build_first_header_map() are supported.

    Returns:

        row_number,
        header_map
    """

    maximum_row = min(
        worksheet.max_row or 1,
        max_search_rows,
    )

    normalized_required_headers = {
        normalize_header(
            header
        )
        for header
        in required_headers
    }

    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=1,
            max_row=maximum_row,
            values_only=True,
        ),
        start=1,
    ):

        header_map = (
            build_first_header_map(
                row
            )
        )

        if (
            normalized_required_headers
            .issubset(
                header_map.keys()
            )
        ):
            return (
                row_number,
                header_map,
            )

    return (
        None,
        {},
    )


def get_row_value(
    row,
    header_map,
    header,
):
    """
    Safely obtain one value from a tuple
    returned by openpyxl iter_rows().
    """

    index = header_map.get(
        normalize_header(
            header
        )
    )

    if index is None:
        return None

    if index >= len(row):
        return None

    return row[index]


def get_first_row_value(
    row,
    header_map,
    *headers,
):
    """
    Return the first meaningful value from
    a list of possible DA6 header names.
    """

    for header in headers:

        value = get_row_value(
            row,
            header_map,
            header,
        )

        if meaningful(
            value
        ):
            return value

    return None


def normalize_level_code(value):
    """
    Examples:

    1  -> 01
    2  -> 02
    10 -> 10
    B1 -> B1
    """

    code = clean_text(
        value
    )

    if not code:
        return ""

    code = code.upper()

    if code.isdigit():
        code = code.zfill(
            2
        )

    return code


def normalize_room_code(value):
    """
    Examples:

    1   -> 001
    5   -> 005
    25  -> 025
    """

    code = clean_text(
        value
    )

    if not code:
        return ""

    if code.isdigit():
        code = code.zfill(
            3
        )

    return code


def decimal_value(
    value,
    field_name,
    row_errors,
):
    if value in (
        None,
        "",
        "-",
    ):
        return None

    try:

        result = Decimal(
            str(value)
        )

        if result < 0:

            row_errors.append(
                (
                    f"{field_name} "
                    "cannot be negative."
                )
            )

        return result

    except (
        InvalidOperation,
        ValueError,
        TypeError,
    ):

        row_errors.append(
            (
                f"{field_name} "
                "must be a number."
            )
        )

        return None


def parse_excel_date(
    value,
    field_name,
    row_errors,
):
    if value in (
        None,
        "",
        "-",
    ):
        return None

    if isinstance(
        value,
        datetime,
    ):
        return value.date()

    if isinstance(
        value,
        date,
    ):
        return value

    text = clean_text(
        value
    )

    date_formats = [
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%d.%m.%Y",
        "%d/%m/%y",
    ]

    for date_format in (
        date_formats
    ):

        try:

            return datetime.strptime(
                text,
                date_format,
            ).date()

        except ValueError:
            continue

    row_errors.append(
        (
            f"{field_name} "
            "has an invalid date."
        )
    )

    return None


# =========================================================
# READ PREMISE FROM WORKBOOK
# =========================================================

def read_premise_from_workbook(
    file_path,
):

    errors = []

    try:

        workbook = (
            open_excel_workbook(
                file_path
            )
        )

    except (
        InvalidFileException,
        OSError,
        ValueError,
        KeyError,
    ) as error:

        return {
            "valid": False,
            "errors": [
                (
                    "Unable to read "
                    f"Excel workbook: {error}"
                )
            ],
            "premise_data": None,
        }

    if (
        "Premise"
        not in workbook.sheetnames
    ):

        workbook.close()

        return {
            "valid": False,
            "errors": [
                (
                    "The workbook must contain "
                    "a sheet named 'Premise'."
                )
            ],
            "premise_data": None,
        }

    worksheet = workbook[
        "Premise"
    ]

    header_row, header_map = (
        find_required_header_row(
            worksheet,
            PREMISE_REQUIRED_HEADERS,
        )
    )

    if header_row is None:

        workbook.close()

        return {
            "valid": False,
            "errors": [
                (
                    "The Premise sheet must contain "
                    "DPA NUMBER, PREMISE NAME, "
                    "ADDRESS and STATE."
                )
            ],
            "premise_data": None,
        }

    data_row = None

    for row in (
        worksheet.iter_rows(
            min_row=(
                header_row + 1
            ),
            values_only=True,
        )
    ):

        dpa_number = clean_text(
            get_row_value(
                row,
                header_map,
                "DPA NUMBER",
            )
        )

        premise_name = clean_text(
            get_row_value(
                row,
                header_map,
                "PREMISE NAME",
            )
        )

        if (
            meaningful(
                dpa_number
            )
            or meaningful(
                premise_name
            )
        ):

            data_row = row

            break

    if data_row is None:

        workbook.close()

        return {
            "valid": False,
            "errors": [
                (
                    "No Premise information "
                    "was found in the Premise sheet."
                )
            ],
            "premise_data": None,
        }

    dpa_number = clean_text(
        get_row_value(
            data_row,
            header_map,
            "DPA NUMBER",
        )
    )

    name = clean_text(
        get_row_value(
            data_row,
            header_map,
            "PREMISE NAME",
        )
    )

    address = clean_text(
        get_row_value(
            data_row,
            header_map,
            "ADDRESS",
        )
    )

    state = clean_text(
        get_row_value(
            data_row,
            header_map,
            "STATE",
        )
    )

    ministry = clean_text(
        get_row_value(
            data_row,
            header_map,
            "MINISTRY",
        )
    )

    department = clean_text(
        get_row_value(
            data_row,
            header_map,
            "DEPARTMENT",
        )
    )

    district = clean_text(
        get_row_value(
            data_row,
            header_map,
            "DISTRICT",
        )
    )

    if not dpa_number:

        errors.append(
            "DPA NUMBER is required."
        )

    if not name:

        errors.append(
            "PREMISE NAME is required."
        )

    if not address:

        errors.append(
            "ADDRESS is required."
        )

    if not state:

        errors.append(
            "STATE is required."
        )

    workbook.close()

    if errors:

        return {
            "valid": False,
            "errors": errors,
            "premise_data": None,
        }

    return {
        "valid": True,
        "errors": [],
        "premise_data": {
            "dpa_number": (
                dpa_number
            ),
            "name": (
                name
            ),
            "address": (
                address
            ),
            "ministry": (
                ministry
            ),
            "department": (
                department
            ),
            "state": (
                state
            ),
            "district": (
                district
            ),
        },
    }


# =========================================================
# GET OR CREATE PREMISE FROM WORKBOOK
# =========================================================

def get_or_create_premise_from_workbook(
    file_path,
):

    result = (
        read_premise_from_workbook(
            file_path
        )
    )

    if not result[
        "valid"
    ]:

        return {
            "success": False,
            "premise": None,
            "created": False,
            "warnings": [],
            "errors": (
                result[
                    "errors"
                ]
            ),
        }

    premise_data = (
        result[
            "premise_data"
        ]
    )

    dpa_number = (
        premise_data[
            "dpa_number"
        ]
    )

    # =====================================================
    # CHECK EXISTING DPA
    # =====================================================

    existing_premise = (
        Premise.objects.filter(
            dpa_number__iexact=(
                dpa_number
            )
        ).first()
    )

    warnings = []

    if existing_premise:

        excel_name = (
            premise_data[
                "name"
            ]
            .strip()
            .lower()
        )

        existing_name = (
            existing_premise
            .name
            .strip()
            .lower()
        )

        if (
            excel_name
            != existing_name
        ):

            warnings.append(
                (
                    "The DPA Number already "
                    "exists, but the Premise "
                    "Name in Excel is different. "
                    f"Existing: "
                    f"'{existing_premise.name}'. "
                    f"Excel: "
                    f"'{premise_data['name']}'. "
                    "The existing Premise was "
                    "reused and was not changed."
                )
            )

        return {
            "success": True,
            "premise": (
                existing_premise
            ),
            "created": False,
            "warnings": (
                warnings
            ),
            "errors": [],
        }

    # =====================================================
    # CREATE NEW PREMISE
    # =====================================================

    try:

        with transaction.atomic():

            premise = (
                Premise.objects.create(
                    dpa_number=(
                        premise_data[
                            "dpa_number"
                        ]
                    ),
                    name=(
                        premise_data[
                            "name"
                        ]
                    ),
                    address=(
                        premise_data[
                            "address"
                        ]
                    ),
                    ministry=(
                        premise_data[
                            "ministry"
                        ]
                    ),
                    department=(
                        premise_data[
                            "department"
                        ]
                    ),
                    state=(
                        premise_data[
                            "state"
                        ]
                    ),
                    district=(
                        premise_data[
                            "district"
                        ]
                    ),
                    status="active",
                )
            )

    except Exception as error:

        return {
            "success": False,
            "premise": None,
            "created": False,
            "warnings": [],
            "errors": [
                (
                    "Unable to create "
                    "Premise: "
                    f"{error}"
                )
            ],
        }

    return {
        "success": True,
        "premise": premise,
        "created": True,
        "warnings": [],
        "errors": [],
    }


# =========================================================
# OLD / INDIVIDUAL EXCEL VALIDATION
# =========================================================

def validate_excel_file(
    file_path,
    import_type,
    import_sheet,
):

    errors = []
    total_rows = 0

    try:

        workbook = (
            open_excel_workbook(
                file_path
            )
        )

    except (
        InvalidFileException,
        OSError,
        ValueError,
        KeyError,
    ) as error:

        return {
            "valid": False,
            "errors": [
                (
                    "The uploaded Excel "
                    "file could not be read. "
                    f"{error}"
                )
            ],
            "total_rows": 0,
        }

    if (
        import_sheet
        not in workbook.sheetnames
    ):

        workbook.close()

        return {
            "valid": False,
            "errors": [
                (
                    f"Sheet "
                    f"'{import_sheet}' "
                    "does not exist."
                )
            ],
            "total_rows": 0,
        }

    worksheet = workbook[
        import_sheet
    ]

    # =====================================================
    # ROOMS
    # =====================================================

    if import_type == "rooms":

        header_row, _ = (
            find_header_row(
                worksheet,
                ROOM_HEADERS,
            )
        )

        if header_row is None:

            errors.append(
                (
                    "DAKRuang headers "
                    "could not be found."
                )
            )

        else:

            for row in (
                worksheet.iter_rows(
                    min_row=(
                        header_row + 1
                    ),
                    values_only=True,
                )
            ):

                if any(
                    meaningful(
                        value
                    )
                    for value
                    in row
                ):

                    total_rows += 1

    # =====================================================
    # COMPONENTS
    # =====================================================

    elif import_type == (
        "components"
    ):

        header_row, header_map = (
            find_required_header_row(
                worksheet,
                COMPONENT_REQUIRED_HEADERS,
            )
        )

        if header_row is None:

            errors.append(
                (
                    "DAKKomponen headers "
                    "could not be found."
                )
            )

        else:

            for row in (
                worksheet.iter_rows(
                    min_row=(
                        header_row + 1
                    ),
                    values_only=True,
                )
            ):

                location = (
                    get_row_value(
                        row,
                        header_map,
                        "LOKASI",
                    )
                )

                name = (
                    get_row_value(
                        row,
                        header_map,
                        "NAMA KOMPONEN",
                    )
                )

                component_id = (
                    get_row_value(
                        row,
                        header_map,
                        "NO. PEROLEHAN",
                    )
                )

                if any(
                    meaningful(
                        value
                    )
                    for value
                    in [
                        location,
                        name,
                        component_id,
                    ]
                ):

                    total_rows += 1

    else:

        errors.append(
            "Invalid import type."
        )

    workbook.close()

    return {
        "valid": (
            len(errors) == 0
        ),
        "errors": errors,
        "total_rows": (
            total_rows
        ),
    }


# =========================================================
# DAK BLOK VALIDATION
# =========================================================

def validate_da6_block_rows(
    file_path,
    premise,
):

    errors = []
    valid_rows = []
    total_rows = 0

    try:

        workbook = (
            open_excel_workbook(
                file_path
            )
        )

    except (
        InvalidFileException,
        OSError,
        ValueError,
        KeyError,
    ) as error:

        return {
            "valid": False,
            "errors": [
                (
                    "Unable to read "
                    f"DA6 file: {error}"
                )
            ],
            "valid_rows": [],
            "total_rows": 0,
        }

    if (
        "DAK Blok"
        not in workbook.sheetnames
    ):

        workbook.close()

        return {
            "valid": False,
            "errors": [
                (
                    "DAK Blok sheet "
                    "was not found."
                )
            ],
            "valid_rows": [],
            "total_rows": 0,
        }

    worksheet = workbook[
        "DAK Blok"
    ]

    required_headers = {
        "KOD BLOK / BINAAN LUAR",
        "NAMA BLOK",
    }

    header_row, header_map = (
        find_required_header_row(
            worksheet,
            required_headers,
        )
    )

    if header_row is None:

        workbook.close()

        return {
            "valid": False,
            "errors": [
                (
                    "The DAK Blok "
                    "column headers "
                    "could not be found."
                )
            ],
            "valid_rows": [],
            "total_rows": 0,
        }

    excel_codes = set()

    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=(
                header_row + 1
            ),
            values_only=True,
        ),
        start=(
            header_row + 1
        ),
    ):

        block_code = clean_text(
            get_row_value(
                row,
                header_map,
                "KOD BLOK / BINAAN LUAR",
            )
        )

        block_name = clean_text(
            get_row_value(
                row,
                header_map,
                "NAMA BLOK",
            )
        )

        function = clean_text(
            get_row_value(
                row,
                header_map,
                "FUNGSI BINAAN",
            )
        )

        status_value = clean_text(
            get_row_value(
                row,
                header_map,
                "STATUS BLOK",
            )
        ).lower()

        if (
            not meaningful(
                block_code
            )
            and not meaningful(
                block_name
            )
        ):
            continue

        total_rows += 1

        row_errors = []

        if not meaningful(
            block_code
        ):

            row_errors.append(
                (
                    "Block code "
                    "is required."
                )
            )

        if not meaningful(
            block_name
        ):

            row_errors.append(
                (
                    "Block name "
                    "is required."
                )
            )

        normalized_code = (
            block_code.upper()
        )

        if meaningful(
            normalized_code
        ):

            if (
                normalized_code
                in excel_codes
            ):

                row_errors.append(
                    (
                        "Duplicate block "
                        f"code '{block_code}' "
                        "in DA6."
                    )
                )

            else:

                excel_codes.add(
                    normalized_code
                )

        existing_block = None

        if meaningful(
            block_code
        ):

            existing_block = (
                Block.objects.filter(
                    premise=premise,
                    code__iexact=(
                        block_code
                    ),
                ).first()
            )

        if status_value in {
            "inactive",
            "tidak aktif",
        }:

            status = (
                "inactive"
            )

        else:

            status = (
                "active"
            )

        if row_errors:

            for error in (
                row_errors
            ):

                errors.append(
                    (
                        f"Row "
                        f"{row_number}: "
                        f"{error}"
                    )
                )

            continue

        valid_rows.append(
            {
                "row_number": (
                    row_number
                ),
                "code": (
                    block_code
                ),
                "name": (
                    block_name
                ),
                "function": (
                    function
                ),
                "status": (
                    status
                ),
                "existing_block": (
                    existing_block
                ),
            }
        )

    workbook.close()

    return {
        "valid": (
            len(errors) == 0
        ),
        "errors": (
            errors
        ),
        "valid_rows": (
            valid_rows
        ),
        "total_rows": (
            total_rows
        ),
    }


# =========================================================
# IMPORT DAK BLOK
# =========================================================

def import_da6_blocks(
    file_path,
    premise,
):

    validation = (
        validate_da6_block_rows(
            file_path,
            premise,
        )
    )

    if not validation[
        "valid"
    ]:

        return {
            "success": False,
            "created": 0,
            "updated": 0,
            "skipped": 0,
            "errors": (
                validation[
                    "errors"
                ]
            ),
        }

    created_count = 0
    skipped_count = 0
    errors = []

    for row in (
        validation[
            "valid_rows"
        ]
    ):

        try:

            existing_block = (
                row[
                    "existing_block"
                ]
            )

            if existing_block:

                skipped_count += 1

                continue

            Block.objects.create(
                premise=(
                    premise
                ),
                code=(
                    row["code"]
                ),
                name=(
                    row["name"]
                ),
                function=(
                    row[
                        "function"
                    ]
                ),
                status=(
                    row["status"]
                ),
                structure_type=(
                    "building"
                ),
                security_level=(
                    "staff"
                ),
            )

            created_count += 1

        except Exception as error:

            errors.append(
                (
                    f"Row "
                    f"{row['row_number']}: "
                    f"{error}"
                )
            )

    return {
        "success": (
            len(errors) == 0
        ),
        "created": (
            created_count
        ),
        "updated": 0,
        "skipped": (
            skipped_count
        ),
        "errors": (
            errors
        ),
    }


# =========================================================
# DAKRUANG VALIDATION
# =========================================================

def validate_room_rows(
    file_path,
    import_sheet="DAKRuang",
    premise=None,
):

    errors = []
    warnings = []
    valid_rows = []
    total_rows = 0

    try:

        workbook = (
            open_excel_workbook(
                file_path
            )
        )

    except (
        InvalidFileException,
        OSError,
        ValueError,
        KeyError,
    ) as error:

        return {
            "valid": False,
            "errors": [
                (
                    "Unable to read "
                    f"DA6 file: {error}"
                )
            ],
            "warnings": [],
            "valid_rows": [],
            "total_rows": 0,
        }

    if (
        import_sheet
        not in workbook.sheetnames
    ):

        workbook.close()

        return {
            "valid": False,
            "errors": [
                (
                    f"Sheet "
                    f"'{import_sheet}' "
                    "does not exist."
                )
            ],
            "warnings": [],
            "valid_rows": [],
            "total_rows": 0,
        }

    worksheet = workbook[
        import_sheet
    ]

    header_row, headers = (
        find_header_row(
            worksheet,
            ROOM_HEADERS,
        )
    )

    if header_row is None:

        workbook.close()

        return {
            "valid": False,
            "errors": [
                (
                    "DAKRuang headers "
                    "could not be found."
                )
            ],
            "warnings": [],
            "valid_rows": [],
            "total_rows": 0,
        }

    header_map = {
        header: index
        for index, header
        in enumerate(
            headers
        )
    }

    excel_room_keys = set()

    for row_number, values in enumerate(
        worksheet.iter_rows(
            min_row=(
                header_row + 1
            ),
            values_only=True,
        ),
        start=(
            header_row + 1
        ),
    ):

        block_code = clean_text(
            get_row_value(
                values,
                header_map,
                "KOD BLOK",
            )
        )

        level_name = clean_text(
            get_row_value(
                values,
                header_map,
                "NAMA ARAS",
            )
        )

        level_code = (
            normalize_level_code(
                get_row_value(
                    values,
                    header_map,
                    "KOD ARAS",
                )
            )
        )

        room_name = clean_text(
            get_row_value(
                values,
                header_map,
                "NAMA RUANG",
            )
        )

        room_code = (
            normalize_room_code(
                get_row_value(
                    values,
                    header_map,
                    "KOD RUANG",
                )
            )
        )

        function_name = (
            clean_text(
                get_row_value(
                    values,
                    header_map,
                    "FUNGSI RUANG",
                )
            )
        )

        function_code = (
            clean_text(
                get_row_value(
                    values,
                    header_map,
                    "KOD FUNGSI",
                )
            )
        )

        area_value = (
            get_row_value(
                values,
                header_map,
                "LUAS RUANG (m2)",
            )
        )

        if not any(
            meaningful(
                value
            )
            for value in [
                block_code,
                level_name,
                level_code,
                room_name,
                room_code,
                function_code,
                area_value,
            ]
        ):
            continue

        total_rows += 1

        row_errors = []

        if not block_code:

            row_errors.append(
                (
                    "KOD BLOK "
                    "is required."
                )
            )

        if not level_name:

            row_errors.append(
                (
                    "NAMA ARAS "
                    "is required."
                )
            )

        if not level_code:

            row_errors.append(
                (
                    "KOD ARAS "
                    "is required."
                )
            )

        if not room_name:

            row_errors.append(
                (
                    "NAMA RUANG "
                    "is required."
                )
            )

        if not room_code:

            row_errors.append(
                (
                    "KOD RUANG "
                    "is required."
                )
            )

        block = None

        if block_code:

            block_matches = (
                Block.objects.filter(
                    code__iexact=(
                        block_code
                    )
                )
            )

            if premise is not None:

                block_matches = (
                    block_matches.filter(
                        premise=(
                            premise
                        )
                    )
                )

            block = (
                block_matches.first()
            )

            if block is None:

                row_errors.append(
                    (
                        f"Block "
                        f"'{block_code}' "
                        "does not exist."
                    )
                )

        level = None

        if (
            block
            and level_code
        ):

            level = (
                Level.objects.filter(
                    block=(
                        block
                    ),
                    code__iexact=(
                        level_code
                    ),
                ).first()
            )

        existing_room = None

        if (
            level
            and room_code
        ):

            existing_room = (
                Room.objects.filter(
                    level=(
                        level
                    ),
                    code__iexact=(
                        room_code
                    ),
                ).first()
            )

        room_key = (
            block_code.upper(),
            level_code.upper(),
            room_code.upper(),
        )

        if all(
            room_key
        ):

            if (
                room_key
                in excel_room_keys
            ):

                row_errors.append(
                    (
                        "Duplicate room "
                        "inside DAKRuang: "
                        f"{block_code}."
                        f"{level_code}."
                        f"{room_code}"
                    )
                )

            else:

                excel_room_keys.add(
                    room_key
                )

        room_function = None

        if function_code:

            room_function = (
                RoomFunctionCode
                .objects
                .filter(
                    code__iexact=(
                        function_code
                    ),
                    is_active=True,
                )
                .first()
            )

            if (
                room_function
                is None
            ):

                row_errors.append(
                    (
                        "Room Function Code "
                        f"'{function_code}' "
                        "does not exist."
                    )
                )

            elif (
                function_name
                and room_function.name
            ):

                if (
                    room_function.name
                    .strip()
                    .lower()
                    !=
                    function_name
                    .strip()
                    .lower()
                ):

                    warnings.append(
                        (
                            f"Row "
                            f"{row_number}: "
                            "FUNGSI RUANG "
                            "does not exactly "
                            "match the controlled "
                            "Room Function name."
                        )
                    )

        area = decimal_value(
            area_value,
            "LUAS RUANG (m2)",
            row_errors,
        )

        if row_errors:

            for error in (
                row_errors
            ):

                errors.append(
                    (
                        f"Row "
                        f"{row_number}: "
                        f"{error}"
                    )
                )

            continue

        valid_rows.append(
            {
                "row_number": (
                    row_number
                ),
                "block": (
                    block
                ),
                "block_code": (
                    block_code
                ),
                "level": (
                    level
                ),
                "level_name": (
                    level_name
                ),
                "level_code": (
                    level_code
                ),
                "existing_room": (
                    existing_room
                ),
                "room_code": (
                    room_code
                ),
                "room_name": (
                    room_name
                ),
                "room_function": (
                    room_function
                ),
                "function_code": (
                    function_code
                ),
                "function_name": (
                    function_name
                ),
                "area": (
                    area
                ),
                "room_tag": (
                    f"{block_code}."
                    f"{level_code}."
                    f"{room_code}"
                ),
            }
        )

    workbook.close()

    return {
        "valid": (
            len(errors) == 0
        ),
        "errors": (
            errors
        ),
        "warnings": (
            warnings
        ),
        "valid_rows": (
            valid_rows
        ),
        "total_rows": (
            total_rows
        ),
    }


# =========================================================
# IMPORT DAKRUANG
# =========================================================

def import_da6_rooms(
    file_path,
    premise,
):

    validation = (
        validate_room_rows(
            file_path,
            "DAKRuang",
            premise=(
                premise
            ),
        )
    )

    if not validation[
        "valid"
    ]:

        return {
            "success": False,
            "levels_created": 0,
            "levels_updated": 0,
            "rooms_created": 0,
            "rooms_updated": 0,
            "rooms_skipped": 0,
            "errors": (
                validation[
                    "errors"
                ]
            ),
        }

    levels_created = 0
    rooms_created = 0
    rooms_skipped = 0

    try:

        with transaction.atomic():

            for row in (
                validation[
                    "valid_rows"
                ]
            ):

                # =========================================
                # GET BLOCK
                # =========================================

                block = (
                    Block.objects.filter(
                        premise=(
                            premise
                        ),
                        code__iexact=(
                            row[
                                "block_code"
                            ]
                        ),
                    ).first()
                )

                if not block:

                    raise ValueError(
                        (
                            f"Row "
                            f"{row['row_number']}: "
                            "Block "
                            f"'{row['block_code']}' "
                            "does not exist."
                        )
                    )

                # =========================================
                # LEVEL
                # =========================================

                level = (
                    Level.objects.filter(
                        block=(
                            block
                        ),
                        code__iexact=(
                            row[
                                "level_code"
                            ]
                        ),
                    ).first()
                )

                if not level:

                    level_code = (
                        row[
                            "level_code"
                        ]
                    )

                    if (
                        level_code
                        .isdigit()
                    ):

                        sequence = int(
                            level_code
                        )

                    else:

                        sequence = (
                            block
                            .levels
                            .count()
                            + 1
                        )

                    level = (
                        Level.objects.create(
                            block=(
                                block
                            ),
                            code=(
                                level_code
                            ),
                            name=(
                                row[
                                    "level_name"
                                ]
                            ),
                            sequence=(
                                sequence
                            ),
                        )
                    )

                    levels_created += 1

                # =========================================
                # ROOM
                # =========================================

                room = (
                    Room.objects.filter(
                        level__block__premise=(
                            premise
                        ),
                        level=(
                            level
                        ),
                        code__iexact=(
                            row[
                                "room_code"
                            ]
                        ),
                    ).first()
                )

                if room:

                    rooms_skipped += 1

                    continue

                Room.objects.create(
                    level=(
                        level
                    ),
                    code=(
                        row[
                            "room_code"
                        ]
                    ),
                    name=(
                        row[
                            "room_name"
                        ]
                    ),
                    room_function=(
                        row[
                            "room_function"
                        ]
                    ),
                    function_code=(
                        row[
                            "function_code"
                        ]
                    ),
                    function_name=(
                        row[
                            "function_name"
                        ]
                    ),
                    area=(
                        row[
                            "area"
                        ]
                    ),
                    space_type=(
                        "room"
                    ),
                    security_level=(
                        "staff"
                    ),
                )

                rooms_created += 1

    except Exception as error:

        return {
            "success": False,
            "levels_created": 0,
            "levels_updated": 0,
            "rooms_created": 0,
            "rooms_updated": 0,
            "rooms_skipped": 0,
            "errors": [
                str(
                    error
                )
            ],
        }

    return {
        "success": True,
        "levels_created": (
            levels_created
        ),
        "levels_updated": 0,
        "rooms_created": (
            rooms_created
        ),
        "rooms_updated": 0,
        "rooms_skipped": (
            rooms_skipped
        ),
        "errors": [],
    }


# =========================================================
# DAKKOMPONEN VALIDATION
# =========================================================

def validate_da6_component_rows(
    file_path,
    premise,
):

    errors = []
    valid_rows = []
    total_rows = 0

    try:

        workbook = (
            open_excel_workbook(
                file_path
            )
        )

    except (
        InvalidFileException,
        OSError,
        ValueError,
        KeyError,
    ) as error:

        return {
            "valid": False,
            "errors": [
                (
                    "Unable to read "
                    f"DA6 file: {error}"
                )
            ],
            "valid_rows": [],
            "total_rows": 0,
        }

    # =====================================================
    # CHECK DAKKOMPONEN SHEET
    # =====================================================

    if (
        "DAKKomponen"
        not in workbook.sheetnames
    ):

        workbook.close()

        return {
            "valid": False,
            "errors": [
                (
                    "DAKKomponen sheet "
                    "was not found."
                )
            ],
            "valid_rows": [],
            "total_rows": 0,
        }

    worksheet = workbook[
        "DAKKomponen"
    ]

    # =====================================================
    # FIND HEADER ROW
    # =====================================================

    header_row, header_map = (
        find_required_header_row(
            worksheet,
            COMPONENT_REQUIRED_HEADERS,
        )
    )

    if header_row is None:

        workbook.close()

        return {
            "valid": False,
            "errors": [
                (
                    "The DAKKomponen "
                    "column headers could "
                    "not be found. Required: "
                    "LOKASI, NAMA KOMPONEN, "
                    "KOD BIDANG KEJURUTERAAN, "
                    "SISTEM and NO. PEROLEHAN."
                )
            ],
            "valid_rows": [],
            "total_rows": 0,
        }

    excel_component_ids = set()

    # =====================================================
    # READ COMPONENT ROWS
    # =====================================================

    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=(
                header_row + 1
            ),
            values_only=True,
        ),
        start=(
            header_row + 1
        ),
    ):

        # -------------------------------------------------
        # LOCATION
        # -------------------------------------------------

        location_code = clean_text(
            get_first_row_value(
                row,
                header_map,
                "LOKASI",
                "KOD LOKASI",
            )
        )

        # -------------------------------------------------
        # COMPONENT ID / NO. PEROLEHAN
        # -------------------------------------------------

        component_id = clean_text(
            get_first_row_value(
                row,
                header_map,
                "NO. PEROLEHAN",
                "NO PEROLEHAN",
                "NO ID KOMPONEN",
                "NO. ID KOMPONEN",
            )
        )

        # -------------------------------------------------
        # COMPONENT CODE
        # -------------------------------------------------
        #
        # Some DA6 versions do not have KOD KOMPONEN.
        # In that case NO. PEROLEHAN is used as a safe
        # fallback for the database component_code field.
        #
        # -------------------------------------------------

        component_code = clean_text(
            get_first_row_value(
                row,
                header_map,
                "KOD KOMPONEN",
            )
        )

        if not component_code:

            component_code = (
                component_id
            )

        # -------------------------------------------------
        # COMPONENT NAME
        # -------------------------------------------------

        component_name = clean_text(
            get_row_value(
                row,
                header_map,
                "NAMA KOMPONEN",
            )
        )

        # -------------------------------------------------
        # DISCIPLINE
        # -------------------------------------------------

        discipline = clean_text(
            get_row_value(
                row,
                header_map,
                "KOD BIDANG KEJURUTERAAN",
            )
        ).upper()

        # -------------------------------------------------
        # SYSTEM
        # -------------------------------------------------

        system_code = clean_text(
            get_row_value(
                row,
                header_map,
                "KOD SISTEM",
            )
        )

        system = clean_text(
            get_row_value(
                row,
                header_map,
                "SISTEM",
            )
        )

        # -------------------------------------------------
        # SUBSYSTEM
        # -------------------------------------------------

        subsystem_code = clean_text(
            get_first_row_value(
                row,
                header_map,
                "KOD SUBSISTEM",
                "KOD SUB SISTEM",
            )
        )

        subsystem = clean_text(
            get_first_row_value(
                row,
                header_map,
                "SUB SISTEM",
                "SUBSISTEM",
            )
        )

        # -------------------------------------------------
        # INSTALLATION / ACQUISITION DATE
        # -------------------------------------------------
        #
        # Prefer TARIKH DIPASANG where available.
        # Otherwise use TARIKH PEROLEHAN from DA6.
        #
        # -------------------------------------------------

        installation_value = (
            get_first_row_value(
                row,
                header_map,
                "TARIKH DIPASANG",
                "TARIKH PEROLEHAN",
            )
        )

        # -------------------------------------------------
        # WARRANTY DATE
        # -------------------------------------------------

        warranty_value = (
            get_first_row_value(
                row,
                header_map,
                "TARIKH WARANTI TAMAT",
                "TARIKH TAMAT WARANTI",
            )
        )

        # -------------------------------------------------
        # STATUS
        # -------------------------------------------------

        status_value = clean_text(
            get_first_row_value(
                row,
                header_map,
                "STATUS KOMPONEN",
                "STATUS",
            )
        ).lower()

        # -------------------------------------------------
        # BRAND
        # -------------------------------------------------

        brand = clean_text(
            get_first_row_value(
                row,
                header_map,
                "JENAMA",
                "PEMBUAT",
            )
        )

        # -------------------------------------------------
        # MODEL
        # -------------------------------------------------

        model = clean_text(
            get_first_row_value(
                row,
                header_map,
                "MODEL",
                "NO MODEL",
                "NO. MODEL",
            )
        )

        # -------------------------------------------------
        # SERIAL NUMBER
        # -------------------------------------------------

        serial_number = clean_text(
            get_first_row_value(
                row,
                header_map,
                "NO SIRI",
                "NO. SIRI",
                "NOMBOR SIRI",
            )
        )

        # =================================================
        # IGNORE UNUSED TEMPLATE ROWS
        # =================================================

        if not any(
            meaningful(
                value
            )
            for value in [
                location_code,
                component_name,
                component_id,
            ]
        ):
            continue

        total_rows += 1

        row_errors = []

        # =================================================
        # REQUIRED VALUES
        # =================================================

        if not location_code:

            row_errors.append(
                (
                    "LOKASI "
                    "is required."
                )
            )

        # TEMPORARILY DISABLED
        # if not component_id:
        #     row_errors.append(
        #         (
        #             "NO. PEROLEHAN "
        #             "is required."
        #         )
        #     )


        if not discipline:

            row_errors.append(
                (
                    "KOD BIDANG "
                    "KEJURUTERAAN "
                    "is required."
                )
            )

        # =================================================
        # SYSTEM
        # =================================================

        if not system:

            system = (
                system_code
            )

        if not system:

            row_errors.append(
                (
                    "SISTEM or "
                    "KOD SISTEM "
                    "is required."
                )
            )

        # =================================================
        # SUBSYSTEM
        # =================================================

        if (
            not subsystem
            and subsystem_code
        ):

            subsystem = (
                subsystem_code
            )

        # =================================================
        # DISCIPLINE VALIDATION
        # =================================================

        valid_disciplines = {
            "A",
            "E",
            "M",
            "T",
            "B",
            "L",
        }

        if (
            discipline
            and discipline
            not in valid_disciplines
        ):

            row_errors.append(
                (
                    "KOD BIDANG "
                    "KEJURUTERAAN "
                    f"'{discipline}' "
                    "is invalid."
                )
            )

        # =================================================
        # LOCATION / ROOM
        # =================================================

        room = None

        if location_code:

            room = (
                Room.objects
                .select_related(
                    "level",
                    "level__block",
                    "level__block__premise",
                )
                .filter(
                    room_tag__iexact=(
                        location_code
                    ),
                    level__block__premise=(
                        premise
                    ),
                )
                .first()
            )

            if room is None:

                row_errors.append(
                    (
                        "LOKASI "
                        f"'{location_code}' "
                        "does not match a "
                        "Room under the "
                        "selected Premise."
                    )
                )

        # =================================================
        # DUPLICATE COMPONENT ID INSIDE EXCEL
        # =================================================

        normalized_component_id = ""

        if component_id:

            normalized_component_id = (
                component_id.upper()
            )

        if normalized_component_id:

            if (
                normalized_component_id
                in excel_component_ids
            ):

                row_errors.append(
                    (
                        "Duplicate "
                        "NO. PEROLEHAN "
                        f"'{component_id}' "
                        "inside DAKKomponen."
                    )
                )

            else:

                excel_component_ids.add(
                    normalized_component_id
                )

        # =================================================
        # EXISTING COMPONENT IN DATABASE
        # =================================================

        existing_component = None

        if component_id:

            existing_component = (
                Component.objects.filter(
                    room=room,
                    component_id__iexact=component_id,
                )
                .first()
            )

        # =================================================
        # INSTALLATION / ACQUISITION DATE
        # =================================================

        installation_date = (
            parse_excel_date(
                installation_value,
                "TARIKH PEROLEHAN",
                row_errors,
            )
        )

        # =================================================
        # WARRANTY EXPIRY
        # =================================================

        warranty_expiry = (
            parse_excel_date(
                warranty_value,
                "TARIKH WARANTI TAMAT",
                row_errors,
            )
        )

        if (
            installation_date
            and warranty_expiry
            and warranty_expiry
            < installation_date
        ):

            row_errors.append(
                (
                    "TARIKH WARANTI "
                    "TAMAT cannot be "
                    "before TARIKH "
                    "PEROLEHAN."
                )
            )

        # =================================================
        # STATUS
        # =================================================

        status_map = {
            "draft": "draft",

            "aktif": "active",
            "active": "active",

            "tidak aktif": "inactive",
            "inactive": "inactive",

            "disahkan": "verified",
            "verified": "verified",

            "submitted": "submitted",
            "returned": "returned",
        }

        status = status_map.get(
            status_value,
            "draft",
        )

        # =================================================
        # ROW ERRORS
        # =================================================

        if row_errors:

            for error in (
                row_errors
            ):

                errors.append(
                    (
                        f"Row "
                        f"{row_number}: "
                        f"{error}"
                    )
                )

            continue

        # =================================================
        # VALID COMPONENT ROW
        # =================================================

        valid_rows.append(
            {
                "row_number": (
                    row_number
                ),

                "room": (
                    room
                ),

                "location_code": (
                    location_code
                ),

                "component_id": (
                    component_id
                ),

                "name": (
                    component_name
                ),

                "component_code": (
                    component_code
                ),

                "discipline": (
                    discipline
                ),

                "system": (
                    system
                ),

                "subsystem": (
                    subsystem
                ),

                "brand": (
                    brand
                ),

                "model": (
                    model
                ),

                "serial_number": (
                    serial_number
                ),

                "installation_date": (
                    installation_date
                ),

                "warranty_expiry": (
                    warranty_expiry
                ),

                "status": (
                    status
                ),

                "existing_component": (
                    existing_component
                ),
            }
        )

    workbook.close()

    return {
        "valid": (
            len(errors) == 0
        ),
        "errors": (
            errors
        ),
        "valid_rows": (
            valid_rows
        ),
        "total_rows": (
            total_rows
        ),
    }

# =========================================================
# IMPORT DAKKOMPONEN
# =========================================================

def import_da6_components(
    file_path,
    premise,
):

    validation = validate_da6_component_rows(
        file_path,
        premise,
    )


    if not validation["valid"]:

        return {
            "success": False,
            "created": 0,
            "updated": 0,
            "skipped": 0,
            "errors": validation["errors"],
        }


    created_count = 0
    skipped_count = 0


    try:

        with transaction.atomic():

            for row in validation["valid_rows"]:


                room = row["room"]

                level = room.level

                block = level.block

                premise = block.premise


                # =========================================
                # CREATE HIERARCHY QR FIRST
                # =========================================

                DynamicQRCode.objects.get_or_create(

                    premise=premise,

                    qr_type="premise",

                    defaults={
                        "qr_id": f"QR-PREMISE-{uuid.uuid4()}",
                    },
                )


                DynamicQRCode.objects.get_or_create(

                    block=block,

                    qr_type="block",

                    defaults={

                        "premise": premise,

                        "qr_id": f"QR-BLOCK-{uuid.uuid4()}",
                    },
                )


                DynamicQRCode.objects.get_or_create(

                    level=level,

                    qr_type="level",

                    defaults={

                        "premise": premise,

                        "block": block,

                        "qr_id": f"QR-LEVEL-{uuid.uuid4()}",
                    },
                )


                DynamicQRCode.objects.get_or_create(

                    room=room,

                    qr_type="room",

                    defaults={

                        "premise": premise,

                        "block": block,

                        "level": level,

                        "qr_id": f"QR-ROOM-{uuid.uuid4()}",
                    },
                )


                # =========================================
                # CHECK COMPONENT DUPLICATE
                # =========================================

                existing_component = (
                    Component.objects.filter(
                        component_id__iexact=row["component_id"],
                    )
                    .first()
                )


                if existing_component:

                    component = existing_component

                    skipped_count += 1


                else:

                    component = Component.objects.create(

                        room=room,

                        component_id=row["component_id"],

                        name=row["name"],

                        component_code=row["component_code"],

                        discipline=row["discipline"],

                        system=row["system"],

                        subsystem=row["subsystem"],

                        brand=row["brand"],

                        model=row["model"],

                        serial_number=row["serial_number"],

                        installation_date=row["installation_date"],

                        warranty_expiry=row["warranty_expiry"],

                        status=row["status"],

                        security_level="staff",
                    )


                    created_count += 1



                # =========================================
                # COMPONENT QR
                # =========================================

                DynamicQRCode.objects.get_or_create(

                    component=component,

                    qr_type="component",

                    defaults={

                        "premise": premise,

                        "block": block,

                        "level": level,

                        "room": room,

                        "qr_id": f"QR-COMPONENT-{uuid.uuid4()}",
                    },
                )



    except Exception as error:

        return {

            "success": False,

            "created": 0,

            "updated": 0,

            "skipped": 0,

            "errors":[
                str(error)
            ],

        }


    return {

        "success": True,

        "created": created_count,

        "updated": 0,

        "skipped": skipped_count,

        "errors": [],

    }

# =========================================================
# FULL DA6 WORKBOOK VALIDATION
# =========================================================

def validate_full_da6_file(
    file_path,
):

    errors = []

    # =====================================================
    # REQUIRED WORKSHEETS
    # =====================================================

    required_sheets = [
        "Premise",
        "DAK Blok",
        "DAKRuang",
        "DAKKomponen",
    ]

    try:

        workbook = (
            open_excel_workbook(
                file_path
            )
        )

    except (
        InvalidFileException,
        OSError,
        ValueError,
        KeyError,
    ) as error:

        return {
            "valid": False,
            "errors": [
                (
                    "Unable to read "
                    "DA6 Excel file: "
                    f"{error}"
                )
            ],
            "sheets_found": [],
            "sheet_counts": {},
            "total_rows": 0,
        }

    sheets_found = (
        workbook.sheetnames
    )

    # =====================================================
    # CHECK REQUIRED WORKSHEETS
    # =====================================================

    for sheet_name in (
        required_sheets
    ):

        if (
            sheet_name
            not in sheets_found
        ):

            errors.append(
                (
                    "Required DA6 sheet "
                    f"'{sheet_name}' "
                    "was not found."
                )
            )

    # =====================================================
    # ASSET COUNTS
    # =====================================================

    sheet_counts = {
        "DAK Blok": 0,
        "DAKRuang": 0,
        "DAKKomponen": 0,
    }

    if errors:

        workbook.close()

        return {
            "valid": False,
            "errors": (
                errors
            ),
            "sheets_found": (
                sheets_found
            ),
            "sheet_counts": (
                sheet_counts
            ),
            "total_rows": 0,
        }

    # =====================================================
    # VALIDATE PREMISE SHEET
    # =====================================================

    premise_worksheet = workbook[
        "Premise"
    ]

    premise_header_row, (
        premise_header_map
    ) = find_required_header_row(
        premise_worksheet,
        PREMISE_REQUIRED_HEADERS,
    )

    if (
        premise_header_row
        is None
    ):

        errors.append(
            (
                "The Premise sheet must "
                "contain DPA NUMBER, "
                "PREMISE NAME, ADDRESS "
                "and STATE."
            )
        )

    else:

        premise_data_row = None

        for row in (
            premise_worksheet
            .iter_rows(
                min_row=(
                    premise_header_row
                    + 1
                ),
                values_only=True,
            )
        ):

            dpa_number = clean_text(
                get_row_value(
                    row,
                    premise_header_map,
                    "DPA NUMBER",
                )
            )

            premise_name = clean_text(
                get_row_value(
                    row,
                    premise_header_map,
                    "PREMISE NAME",
                )
            )

            if (
                meaningful(
                    dpa_number
                )
                or meaningful(
                    premise_name
                )
            ):

                premise_data_row = row

                break

        if premise_data_row is None:

            errors.append(
                (
                    "No Premise information "
                    "was found in the "
                    "Premise sheet."
                )
            )

        else:

            dpa_number = clean_text(
                get_row_value(
                    premise_data_row,
                    premise_header_map,
                    "DPA NUMBER",
                )
            )

            premise_name = clean_text(
                get_row_value(
                    premise_data_row,
                    premise_header_map,
                    "PREMISE NAME",
                )
            )

            address = clean_text(
                get_row_value(
                    premise_data_row,
                    premise_header_map,
                    "ADDRESS",
                )
            )

            state = clean_text(
                get_row_value(
                    premise_data_row,
                    premise_header_map,
                    "STATE",
                )
            )

            if not dpa_number:

                errors.append(
                    (
                        "Premise sheet: "
                        "DPA NUMBER is required."
                    )
                )

            if not premise_name:

                errors.append(
                    (
                        "Premise sheet: "
                        "PREMISE NAME is required."
                    )
                )

            if not address:

                errors.append(
                    (
                        "Premise sheet: "
                        "ADDRESS is required."
                    )
                )

            if not state:

                errors.append(
                    (
                        "Premise sheet: "
                        "STATE is required."
                    )
                )

    # =====================================================
    # COUNT DAK BLOK
    # =====================================================

    worksheet = workbook[
        "DAK Blok"
    ]

    header_row, header_map = (
        find_required_header_row(
            worksheet,
            {
                "KOD BLOK / BINAAN LUAR",
                "NAMA BLOK",
            },
        )
    )

    if header_row is not None:

        for row in (
            worksheet.iter_rows(
                min_row=(
                    header_row + 1
                ),
                values_only=True,
            )
        ):

            code = (
                get_row_value(
                    row,
                    header_map,
                    "KOD BLOK / BINAAN LUAR",
                )
            )

            name = (
                get_row_value(
                    row,
                    header_map,
                    "NAMA BLOK",
                )
            )

            if any(
                meaningful(
                    value
                )
                for value in [
                    code,
                    name,
                ]
            ):

                sheet_counts[
                    "DAK Blok"
                ] += 1

    else:

        errors.append(
            (
                "The required DAK Blok "
                "column headers could "
                "not be found."
            )
        )

    # =====================================================
    # COUNT DAKRUANG
    # =====================================================

    worksheet = workbook[
        "DAKRuang"
    ]

    room_header_row, (
        room_headers
    ) = find_header_row(
        worksheet,
        ROOM_HEADERS,
    )

    if (
        room_header_row
        is not None
    ):

        room_header_map = {
            header: index
            for index, header
            in enumerate(
                room_headers
            )
        }

        for row in (
            worksheet.iter_rows(
                min_row=(
                    room_header_row
                    + 1
                ),
                values_only=True,
            )
        ):

            block_code = (
                get_row_value(
                    row,
                    room_header_map,
                    "KOD BLOK",
                )
            )

            level_code = (
                get_row_value(
                    row,
                    room_header_map,
                    "KOD ARAS",
                )
            )

            room_code = (
                get_row_value(
                    row,
                    room_header_map,
                    "KOD RUANG",
                )
            )

            if (
                meaningful(
                    block_code
                )
                and meaningful(
                    level_code
                )
                and meaningful(
                    room_code
                )
            ):

                sheet_counts[
                    "DAKRuang"
                ] += 1

    else:

        errors.append(
            (
                "The required DAKRuang "
                "column headers could "
                "not be found."
            )
        )

    # =====================================================
    # COUNT DAKKOMPONEN
    # =====================================================

    worksheet = workbook[
        "DAKKomponen"
    ]

    component_header_row, (
        component_header_map
    ) = find_required_header_row(
        worksheet,
        COMPONENT_REQUIRED_HEADERS,
    )

    if (
        component_header_row
        is not None
    ):

        for row in (
            worksheet.iter_rows(
                min_row=(
                    component_header_row
                    + 1
                ),
                values_only=True,
            )
        ):

            location = (
                get_first_row_value(
                    row,
                    component_header_map,
                    "LOKASI",
                    "KOD LOKASI",
                )
            )

            component_name = (
                get_row_value(
                    row,
                    component_header_map,
                    "NAMA KOMPONEN",
                )
            )

            component_id = (
                get_first_row_value(
                    row,
                    component_header_map,
                    "NO. PEROLEHAN",
                    "NO PEROLEHAN",
                    "NO ID KOMPONEN",
                    "NO. ID KOMPONEN",
                )
            )

            if any(
                meaningful(
                    value
                )
                for value in [
                    location,
                    component_name,
                    component_id,
                ]
            ):

                sheet_counts[
                    "DAKKomponen"
                ] += 1

    else:

        errors.append(
            (
                "The required "
                "DAKKomponen column "
                "headers could not "
                "be found."
            )
        )

    # =====================================================
    # CLOSE WORKBOOK
    # =====================================================

    workbook.close()

    # =====================================================
    # TOTAL ASSET ROWS
    # =====================================================

    total_rows = sum(
        sheet_counts.values()
    )

    # =====================================================
    # FINAL RESULT
    # =====================================================

    return {
        "valid": (
            len(errors) == 0
        ),
        "errors": (
            errors
        ),
        "sheets_found": (
            sheets_found
        ),
        "sheet_counts": (
            sheet_counts
        ),
        "total_rows": (
            total_rows
        ),
    }